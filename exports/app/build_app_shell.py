"""Build the themed SHELL (.xlsx) for the live SteelBuild Pro Excel app.
Static layout only — banners, headers, Home/control sheet, hidden _config.
VBA fills & styles all data rows at runtime. Also emits vba/mSchema.bas so the
column definitions are single-sourced here.
Run: python build_app_shell.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(r"C:\dev\SteelBuild-Pro-Rev.2\exports\app")
VBA = BASE / "vba"; VBA.mkdir(parents=True, exist_ok=True)
OUT = BASE / "SteelBuild-Pro-App.xlsx"

SUPABASE_URL = "https://kjrwqagyeswwoxpjkcko.supabase.co"
ANON_KEY = ("eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImtqcndxYWd5"
            "ZXN3d294cGprY2tvIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzU0NjA5OTAsImV4cCI6MjA5MTAzNjk5MH0."
            "2JvQQizgxARmDY-55zlWI18NJ5e1TpHBAKg_0Cw5L58")

CANVAS="0D1117"; PANEL="161B22"; PANEL2="1F2937"; BAND="11161C"
INK="E6EDF3"; SUB="C9D1D9"; MUTE="8B949E"; LINE="30363D"
BLUE="3B82F6"; CYAN="38BDF8"; GREEN="22C55E"; AMBER="F59E0B"; RED="EF4444"
PURPLE="A78BFA"; TEAL="2DD4BF"; SLATE="64748B"; ORANGE="FB923C"

def F(size=10, bold=False, color=INK):
    return Font(name="Arial", size=size, bold=bold, color=color)
def fill(h): return PatternFill("solid", fgColor=h)
LEFT=Alignment("left", vertical="center")
CTR=Alignment("center", vertical="center")
START=2

# --- module/table schema: single source of truth -----------------------
# type codes used by VBA: text date num money tons pct bool check wrap
#   chip:<domain> -> coloured status cell
SCHEMA = [
 dict(key="drawing_sets", table="drawing_sets", sheet="Drawing Sets", pf=True,
      title="Drawing Sets", sub="Packages tracked at the set level",
      order="issued_date.asc", cols=[
        ("set_name","Set Name",26,"text"),("discipline","Discipline",13,"text"),
        ("revision","Rev",6,"text"),("status","Status",10,"chip:sub_status"),
        ("set_approval_status","Approval",16,"chip:set_approval"),
        ("is_locked","Locked",8,"check"),("sheet_count","Sheets",8,"num"),
        ("issued_date","Issued",12,"date"),("locked_reason","Lock Reason",44,"text")]),
 dict(key="drawings", table="drawings", sheet="Drawings", pf=True,
      title="Drawings", sub="Individual sheets with stage & revision",
      order="sheet_number.asc", cols=[
        ("sheet_number","Sheet #",13,"text"),("title","Title",36,"text"),
        ("discipline","Discipline",13,"text"),("revision_number","Rev",6,"text"),
        ("stage","Stage",12,"chip:dwg_stage"),("drawing_set_name","Drawing Set",22,"text"),
        ("submitted_date","Submitted",11,"date"),("return_date","Returned",11,"date"),
        ("is_superseded","Superseded",11,"check")]),
 dict(key="submittals", table="submittals", sheet="Submittals", pf=True,
      title="Submittals", sub="Not Started -> IFA -> OFA -> BFA -> OFS -> IFC -> Released for Fab",
      order="submitted_date.asc", cols=[
        ("submittal_number","Sub #",9,"text"),("title","Title",24,"text"),
        ("spec_section","Spec",10,"text"),("discipline","Discipline",13,"text"),
        ("revision","Rev",6,"text"),("round_number","Rnd",6,"num"),
        ("status","Status",20,"chip:sub_status"),("ball_in_court","Ball in Court",13,"text"),
        ("reviewer","Reviewer",18,"text"),("submitted_date","Submitted",11,"date"),
        ("required_date","Required",11,"date"),("returned_date","Returned",11,"date")]),
 dict(key="rfis", table="rfis", sheet="RFIs", pf=True,
      title="RFIs", sub="Requests for information with cost & schedule impact",
      order="submitted_date.asc", cols=[
        ("rfi_number","RFI #",10,"text"),("title","Subject",34,"text"),
        ("discipline","Discipline",11,"text"),("status","Status",17,"chip:rfi_status"),
        ("priority","Priority",9,"chip:priority"),("ball_in_court","Ball",9,"text"),
        ("assigned_to","Assigned",11,"text"),("submitted_date","Submitted",11,"date"),
        ("date_answered","Answered",11,"date"),("cost_impact","Cost",6,"check"),
        ("schedule_impact","Sched",6,"check"),("question","Question",40,"wrap"),
        ("answer","Answer",40,"wrap")]),
 dict(key="schedule_tasks", table="schedule_tasks", sheet="Schedule", pf=True,
      title="Schedule", sub="WBS tasks across the project phases",
      order="start_date.asc", cols=[
        ("wbs_code","WBS",7,"text"),("task_name","Task",34,"text"),
        ("task_type","Type",12,"text"),("phase","Phase",15,"text"),
        ("status","Status",13,"chip:sched_status"),("priority","Priority",9,"chip:priority"),
        ("percent_complete","% Complete",14,"pct"),("start_date","Start",11,"date"),
        ("end_date","Finish",11,"date"),("crew_name","Crew",12,"text"),
        ("is_milestone","Milestone",9,"check")]),
 dict(key="work_packages", table="work_packages", sheet="Work Packages", pf=True,
      title="Work Packages", sub="Released scope by area & crew",
      order="wp_number.asc", cols=[
        ("wp_number","WP #",9,"text"),("name","Package",28,"text"),
        ("phase","Phase",13,"text"),("status","Status",13,"chip:wp_status"),
        ("percent_complete","% Complete",14,"pct"),("area","Area",15,"text"),
        ("crew","Crew",12,"text"),("released_date","Released",12,"date"),
        ("scheduled_start_date","Start",11,"date"),("scheduled_end_date","Finish",11,"date")]),
 dict(key="deliveries", table="deliveries", sheet="Deliveries", pf=True,
      title="Deliveries", sub="Truckloads, pieces & tonnage",
      order="scheduled_date.asc", cols=[
        ("load_number","Load",8,"text"),("description","Description",26,"text"),
        ("vendor","Vendor",10,"text"),("status","Status",11,"chip:sub_status"),
        ("priority","Priority",9,"chip:priority"),("pieces","Pieces",9,"num"),
        ("weight_tons","Tons",9,"tons"),("area","Area",12,"text"),
        ("scheduled_date","Scheduled",12,"date"),("actual_date","Delivered",12,"date")]),
 dict(key="change_orders", table="change_orders", sheet="Change Orders", pf=True,
      title="Change Orders", sub="Contract changes & margin",
      order="co_number.asc", cols=[
        ("co_number","CO #",9,"text"),("title","Title",28,"text"),
        ("status","Status",12,"chip:co_status"),("co_amount","Amount",13,"money"),
        ("reason_code","Reason",15,"text"),("schedule_impact_days","Sched Days",11,"num"),
        ("submitted_date","Submitted",12,"date"),("approved_date","Approved",12,"date"),
        ("description","Description",46,"wrap")]),
 dict(key="vendors", table="vendors", sheet="Vendors", pf=False,
      title="Vendors", sub="Subcontractors, suppliers & service providers",
      order="company_name.asc", cols=[
        ("company_name","Company",20,"text"),("vendor_type","Type",15,"text"),
        ("contact_person","Contact",18,"text"),("title","Title",14,"text"),
        ("phone","Phone",15,"text"),("email","Email",30,"text"),
        ("state","State",7,"text"),("is_preferred","Preferred",10,"check"),
        ("payment_terms","Terms",10,"text"),("status","Status",9,"chip:sub_status")]),
]
HEADER_ROW = 6  # every module table header sits on row 6

wb = Workbook()
# default font -> Arial
try:
    wb._named_styles['Normal'].font = Font(name="Arial", size=10, color="000000")
except Exception:
    pass

def add_name(nm, ref):
    wb.defined_names.add(DefinedName(nm, attr_text=ref))

def paint(ws, cols, rows=320):
    for r in range(1, rows+1):
        for c in range(1, cols+3):
            ws.cell(r,c).fill=fill(CANVAS)
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.tabColor=PANEL2
    ws.column_dimensions[get_column_letter(1)].width=2.4

def banner(ws, last_col, title, subtitle):
    for c in range(START,last_col+1): ws.cell(2,c).fill=fill(PANEL2); ws.cell(3,c).fill=fill(PANEL2)
    ws.merge_cells(start_row=2,start_column=START,end_row=2,end_column=last_col)
    ws.merge_cells(start_row=3,start_column=START,end_row=3,end_column=last_col)
    t=ws.cell(2,START,title); t.font=F(20,True,"FFFFFF"); t.alignment=LEFT
    s=ws.cell(3,START,subtitle); s.font=F(10,False,CYAN); s.alignment=LEFT
    for c in range(START,last_col+1): ws.cell(4,c).fill=fill(BLUE)
    ws.row_dimensions[2].height=30; ws.row_dimensions[3].height=18; ws.row_dimensions[4].height=3

def header_row(ws, columns, row):
    for j,(key,label,width,typ) in enumerate(columns):
        c=ws.cell(row, START+j, label)
        c.font=F(9,True,CYAN); c.fill=fill(PANEL2)
        c.alignment=Alignment("left" if typ in ("text","wrap","date") else "center", vertical="center", wrap_text=True)
        c.border=Border(bottom=Side(style="medium", color=BLUE))
        ws.column_dimensions[get_column_letter(START+j)].width=width
    ws.row_dimensions[row].height=24
    ws.freeze_panes=ws.cell(row+1, START)

def section(ws,r,last_col,text,color=CYAN):
    for c in range(START,last_col+1): ws.cell(r,c).fill=fill(BAND)
    ws.merge_cells(start_row=r,start_column=START,end_row=r,end_column=last_col)
    ws.cell(r,START,text).font=F(11,True,color); ws.cell(r,START).alignment=Alignment("left",vertical="center",indent=1)
    ws.row_dimensions[r].height=20

# ---- module sheets -----------------------------------------------------
order_sheets=["Home","Command Center","Detailing Hub"]+[m["sheet"] for m in SCHEMA]+["Fabrication","Financials","About"]
ws0=wb.active; ws0.title="Home"
created={"Home":ws0}
for nm in order_sheets[1:]:
    created[nm]=wb.create_sheet(nm)
cfg=wb.create_sheet("_config"); cfg.sheet_state="hidden"

for m in SCHEMA:
    ws=created[m["sheet"]]; lc=START+len(m["cols"])-1
    paint(ws, lc)
    banner(ws, lc, m["title"], m["sub"])
    ws.cell(5,START,"  Sign in on the Home tab, pick a project, and press Refresh to load live data.").font=F(9,False,MUTE)
    header_row(ws, m["cols"], HEADER_ROW)

# ---- Command Center (KPI + funnel + ball-in-court + exceptions) --------
ws=created["Command Center"]; LC=12
paint(ws,LC); banner(ws,LC,"Command Center","Live operational rollups for the selected project")
KPI=[("kc_sheets","Drawing Sheets",CYAN),("kc_sets","Drawing Sets",BLUE),
     ("kc_subs","Submittals",TEAL),("kc_openrfi","Open RFIs",AMBER),("kc_costrfi","RFIs · Cost Impact",RED),
     ("kc_schrfi","RFIs · Sched Impact",ORANGE),("kc_tons","Tons Delivered",PURPLE),
     ("kc_sched","Schedule % Complete",GREEN),("kc_wps","Work Packages",BLUE),("kc_delay","Delayed Tasks",RED)]
section(ws,6,LC,"KEY PERFORMANCE INDICATORS"); r=7; c=START
for i,(nm,label,acc) in enumerate(KPI):
    for rr in (r,r+1,r+2):
        for cc in range(c,c+2): ws.cell(rr,cc).fill=fill(PANEL)
    for cc in range(c,c+2): ws.cell(r+3,cc).fill=fill(acc)
    ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+1)
    v=ws.cell(r,c,"—"); v.font=F(16,True,INK); v.alignment=Alignment("left",vertical="center",indent=1)
    add_name(nm, f"'Command Center'!{get_column_letter(c)}{r}")
    ws.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+1)
    ws.cell(r+1,c,label.upper()).font=F(8,True,MUTE); ws.cell(r+1,c).alignment=Alignment("left",vertical="center",indent=1)
    ws.row_dimensions[r].height=24; ws.row_dimensions[r+1].height=13; ws.row_dimensions[r+2].height=4; ws.row_dimensions[r+3].height=3
    c+=2
    if i==4: r+=5; c=START
r+=5
section(ws,r,LC,"SUBMITTAL STATUS"); r+=1
add_name("kc_funnel", f"'Command Center'!{get_column_letter(START)}{r}")  # VBA writes 3 rows here
r+=4
section(ws,r,LC//2+1 if False else START+3,"BALL IN COURT");
# two half headers
for c in range(START,START+4): ws.cell(r,c).fill=fill(BAND)
ws.merge_cells(start_row=r,start_column=START,end_row=r,end_column=START+3); ws.cell(r,START,"  BALL IN COURT").font=F(11,True,CYAN)
for c in range(START+5,LC+1): ws.cell(r,c).fill=fill(BAND)
ws.merge_cells(start_row=r,start_column=START+5,end_row=r,end_column=LC); ws.cell(r,START+5,"  EXCEPTIONS").font=F(11,True,AMBER)
ws.row_dimensions[r].height=20
add_name("kc_bic", f"'Command Center'!{get_column_letter(START)}{r+1}")
add_name("kc_exc", f"'Command Center'!{get_column_letter(START+5)}{r+1}")

# ---- Detailing Hub -----------------------------------------------------
ws=created["Detailing Hub"]; LC=10
paint(ws,LC); banner(ws,LC,"Detailing Control Center","Drawing sets <-> submittals (auto-lock) + open RFIs")
hub_cols=[("set_name","Drawing Set",26,"text"),("discipline","Discipline",13,"text"),
          ("revision","Rev",6,"text"),("sheet_count","Sheets",8,"num"),
          ("set_approval_status","Set Approval",16,"chip:set_approval"),
          ("is_locked","Locked",8,"check"),("linked_submittal","Submittal",11,"text"),
          ("submittal_status","Submittal Status",20,"chip:sub_status"),("ball_in_court","Ball in Court",13,"text")]
section(ws,6,LC,"SET <-> SUBMITTAL LINKAGE  (auto-lock fires when the submittal is approved)")
header_row(ws, hub_cols, 7)
add_name("hub_open", "'Detailing Hub'!B100")  # VBA places the open-RFI table dynamically; anchor reserved

# ---- Fabrication -------------------------------------------------------
ws=created["Fabrication"]; LC=10
paint(ws,LC); banner(ws,LC,"Fabrication & Production","Fab-release readiness + production rollup from the schedule")
section(ws,6,LC,"FAB-RELEASE READINESS")
add_name("fab_ready","Fabrication!B7")
section(ws,11,LC,"FABRICATION PROGRESS  ·  schedule (Fabrication phase)")
header_row(ws,[("task_name","Fabrication Activity",34,"text"),("status","Status",13,"chip:sched_status"),
               ("percent_complete","% Complete",14,"pct"),("start_date","Start",11,"date"),
               ("end_date","Finish",11,"date")],12)
add_name("fab_tasks","Fabrication!B13")

# ---- Financials --------------------------------------------------------
ws=created["Financials"]; LC=9
paint(ws,LC); banner(ws,LC,"Financials","Contract value, change orders & AIA G702/G703")
section(ws,6,LC,"CONTRACT SUMMARY")
fin_labels=["Original Contract Value","Net Approved Change Orders","Adjusted Contract Value",
            "Retainage","Contract Less Retainage"]
for i,lab in enumerate(fin_labels):
    r=7+i
    for c in range(START,START+4): ws.cell(r,c).fill=fill(PANEL)
    ws.merge_cells(start_row=r,start_column=START,end_row=r,end_column=START+3)
    ws.cell(r,START,"  "+lab).font=F(10,True,INK)
    for c in range(START+4,START+6): ws.cell(r,c).fill=fill(PANEL)
    ws.merge_cells(start_row=r,start_column=START+4,end_row=r,end_column=START+5)
    cc=ws.cell(r,START+4,"—"); cc.font=F(11,True,CYAN); cc.alignment=Alignment("right",vertical="center"); cc.number_format='$#,##0;($#,##0)'
add_name("fin_orig","Financials!F7"); add_name("fin_netco","Financials!F8")
add_name("fin_adj","Financials!F9"); add_name("fin_ret","Financials!F10"); add_name("fin_less","Financials!F11")
section(ws,13,LC,"CHANGE ORDERS  (this project)")
header_row(ws,[("co_number","CO #",9,"text"),("title","Title",30,"text"),
               ("status","Status",12,"chip:co_status"),("co_amount","Amount",14,"money"),
               ("reason_code","Reason",16,"text"),("approved_date","Approved",12,"date")],14)
add_name("fin_cos","Financials!B15")

# ---- About -------------------------------------------------------------
ws=created["About"]; LC=10
paint(ws,LC); banner(ws,LC,"About this Workbook","Live SteelBuild Pro client · how it works")
about=[("HOW IT WORKS",CYAN,[
   "This workbook is a live, multi-project client for SteelBuild Pro. On the Home tab, Sign In with your",
   "SteelBuild Pro email & password, choose a project, and press Refresh. Data is pulled straight from the",
   "live database over a secure connection and written into every module tab.",
   "Switch projects any time from the Home dropdown and Refresh again."]),
 ("SECURITY",GREEN,[
   "The workbook stores only the PUBLIC anon key (the same one your web app ships). Your password is used",
   "once at sign-in to obtain a short-lived token and is never written to disk. Every read & write is",
   "scoped by row-level security to the projects your account can access. No service key is embedded."]),
 ("WHAT YOU CAN DO",TEAL,[
   "Read every module live · switch projects · live dashboards & exceptions · the Detailing auto-lock view.",
   "Create an RFI, add a Submittal, and advance a submittal's status — each is confirmed before it writes,",
   "and each writes an activity audit row so nothing changes un-logged. Higher-risk bulk edits are not enabled."]),
 ("STATUS COLOURS",PURPLE,[
   "Green = approved / released / complete.  Teal = approved as noted.  Amber = in review / in progress.",
   "Blue = active stage.  Red = delayed / revise & resubmit / cost-or-schedule impact.  Gray = inactive."]),
]
r=6
for title,acc,lines in about:
    section(ws,r,LC,title,acc); r+=1
    for ln in lines:
        for c in range(START,LC+1): ws.cell(r,c).fill=fill(PANEL)
        ws.merge_cells(start_row=r,start_column=START,end_row=r,end_column=LC)
        ws.cell(r,START,"   "+ln).font=F(9,False,SUB); ws.row_dimensions[r].height=15; r+=1
    r+=1

# ---- Home --------------------------------------------------------------
ws=created["Home"]; LC=12
paint(ws,LC)
for c in range(START,LC+1): ws.cell(2,c).fill=fill(PANEL2); ws.cell(3,c).fill=fill(PANEL2)
ws.merge_cells(start_row=2,start_column=START,end_row=2,end_column=LC)
ws.merge_cells(start_row=3,start_column=START,end_row=3,end_column=LC)
ws.cell(2,START,"STEELBUILD  PRO").font=F(24,True,"FFFFFF"); ws.cell(2,START).alignment=LEFT
ws.cell(3,START,"Live Project Workspace  ·  Excel Edition").font=F(11,False,CYAN); ws.cell(3,START).alignment=LEFT
for c in range(START,LC+1): ws.cell(4,c).fill=fill(BLUE)
ws.row_dimensions[2].height=34; ws.row_dimensions[3].height=18; ws.row_dimensions[4].height=3
# connection panel
def kv(r, label, value, vcolor=INK, name=None, vcols=4):
    for c in range(START,START+2): ws.cell(r,c).fill=fill(BAND)
    ws.merge_cells(start_row=r,start_column=START,end_row=r,end_column=START+1)
    ws.cell(r,START,"  "+label).font=F(9,True,MUTE)
    for c in range(START+2,START+2+vcols): ws.cell(r,c).fill=fill(PANEL)
    ws.merge_cells(start_row=r,start_column=START+2,end_row=r,end_column=START+1+vcols)
    cc=ws.cell(r,START+2,value); cc.font=F(11,True,vcolor); cc.alignment=Alignment("left",vertical="center",indent=1)
    if name: add_name(name, f"Home!{get_column_letter(START+2)}{r}")
    ws.row_dimensions[r].height=20
section(ws,6,LC,"CONNECTION")
kv(7,"STATUS","Not signed in",AMBER,"home_status",6)
kv(8,"SIGNED IN AS","—",SUB,"home_user",6)
section(ws,10,LC,"PROJECT")
# project dropdown cell (validation list points at _config)
for c in range(START,START+2): ws.cell(11,c).fill=fill(BAND)
ws.merge_cells(start_row=11,start_column=START,end_row=11,end_column=START+1); ws.cell(11,START,"  SELECT PROJECT").font=F(9,True,MUTE)
for c in range(START+2,START+8): ws.cell(11,c).fill=fill(PANEL)
ws.merge_cells(start_row=11,start_column=START+2,end_row=11,end_column=START+7)
pick=ws.cell(11,START+2,""); pick.font=F(11,True,CYAN); pick.alignment=Alignment("left",vertical="center",indent=1)
add_name("home_project", f"Home!{get_column_letter(START+2)}11")
dv=DataValidation(type="list", formula1="=cfgProjDisplay", allow_blank=True); ws.add_data_validation(dv); dv.add(pick)
kv(12,"LAST REFRESHED","—",SUB,"home_last",6)
section(ws,14,LC,"ACTIONS  ·  buttons added below")
ws.cell(15,START,"   Sign In  /  Refresh  /  Sign Out  ·  Create RFI  ·  Add Submittal  ·  Advance Submittal").font=F(9,False,MUTE)
add_name("home_btnrow","Home!B17")
section(ws,20,LC,"GO TO MODULE")
nav=["Command Center","Detailing Hub","Drawing Sets","Drawings","Submittals","RFIs","Schedule",
     "Work Packages","Deliveries","Fabrication","Change Orders","Financials","Vendors","About"]
r=21
for i,nm in enumerate(nav):
    col=START if i%2==0 else START+6
    for c in range(col,col+5): ws.cell(r,c).fill=fill(PANEL)
    ws.merge_cells(start_row=r,start_column=col,end_row=r,end_column=col+4)
    lk=ws.cell(r,col,f"   ▸  {nm}"); lk.font=F(10,True,CYAN); lk.hyperlink=f"#'{nm}'!A1"; lk.alignment=LEFT
    ws.row_dimensions[r].height=17
    if i%2==1: r+=1
if len(nav)%2==1: r+=1
ws.cell(r+1,START,"SteelBuild Pro live client  ·  reads & writes are RLS-scoped to your account  ·  no service key embedded").font=F(8,True,MUTE)

# ---- _config -----------------------------------------------------------
cfg.sheet_view.showGridLines=False
rows=[("URL",SUPABASE_URL),("ANON",ANON_KEY),("EMAIL",""),("TOKEN",""),("REFRESH_TOKEN",""),
      ("PROJ_ID",""),("PROJ_NUM",""),("PROJ_NAME",""),("STATUS","Not signed in"),("LAST_REFRESH","")]
for i,(k,v) in enumerate(rows):
    cfg.cell(2+i,1,k); cfg.cell(2+i,2,v)
add_name("cfgUrl","_config!$B$2"); add_name("cfgAnon","_config!$B$3"); add_name("cfgEmail","_config!$B$4")
add_name("cfgToken","_config!$B$5"); add_name("cfgRefresh","_config!$B$6")
add_name("cfgProjId","_config!$B$7"); add_name("cfgProjNum","_config!$B$8"); add_name("cfgProjName","_config!$B$9")
add_name("cfgStatus","_config!$B$10"); add_name("cfgLast","_config!$B$11")
# projects list area (VBA fills): A=id B=num C=name D=display, rows 21..200
cfg.cell(20,1,"id"); cfg.cell(20,2,"number"); cfg.cell(20,3,"name"); cfg.cell(20,4,"display")
add_name("cfgProjIds","_config!$A$21:$A$200")
add_name("cfgProjDisplay","_config!$D$21:$D$200")

wb.save(OUT)
print("SHELL SAVED", OUT)

# ---- emit vba/mSchema.bas ---------------------------------------------
def esc(s): return s.replace('"','""')
lines=['Attribute VB_Name = "mSchema"',"'Auto-generated by build_app_shell.py - do not edit by hand","Option Explicit",""]
lines.append("Public Const SB_URL As String = \"%s\"" % SUPABASE_URL)
lines.append("Public Const SB_ANON As String = \"%s\"" % ANON_KEY)
lines.append("Public Const HEADER_ROW As Long = %d" % HEADER_ROW)
lines.append("")
lines.append("' Returns array of module defs; each = Array(key, table, sheet, projectFiltered, order, colsCsv, typesCsv)")
lines.append("Public Function Modules() As Variant")
lines.append("    Dim m(0 To %d) As Variant" % (len(SCHEMA)-1))
for i,mdef in enumerate(SCHEMA):
    keys=",".join(c[0] for c in mdef["cols"])
    types=",".join(c[3] for c in mdef["cols"])
    lines.append('    m(%d) = Array("%s","%s","%s",%s,"%s","%s","%s")' % (
        i, mdef["key"], mdef["table"], esc(mdef["sheet"]),
        "True" if mdef["pf"] else "False", mdef["order"], keys, types))
lines.append("    Modules = m")
lines.append("End Function")
(VBA/"mSchema.bas").write_text("\r\n".join(lines), encoding="ascii")
print("EMITTED", VBA/"mSchema.bas")
